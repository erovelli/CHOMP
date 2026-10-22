#!/usr/bin/env python3
"""Figure 2 as an editable PowerPoint chart: NHE vs. HHS release vs. mapped total.

Same data and styling as scripts/plot_nhe.py, but written as a native PowerPoint
line chart so co-authors can restyle it and edit the numbers (right-click the
chart > Edit Data opens the embedded Excel sheet). The slide is sized to the
figure, so File > Export / Save as PDF gives a submission-ready figure.

The percentage labels under the mapped line are custom data-label text: they stay
attached to their points but do not update if the numbers are edited. Series are
identified by the legend and marker shape (no end-of-line labels, so the chart
stays simple to edit).

Reads docs/validation/nhe_benchmark.csv, writes docs/validation/fig_nhe_benchmark.pptx.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.util import Inches, Pt

REPO_ROOT = Path(__file__).resolve().parent.parent
VAL = REPO_ROOT / "docs" / "validation"

# Same colors and markers as plot_nhe.py. (column, legend label, hex, marker)
SERIES = [
    ("medicaid_chip", "NHE Medicaid + CHIP dental services", "2A78D6", XL_MARKER_STYLE.CIRCLE),
    ("release_dental", "HHS release, all dental codes", "EB6834", XL_MARKER_STYLE.SQUARE),
    ("chomp_50dc", "Mapped to 50 states + DC", "1BAF7A", XL_MARKER_STYLE.TRIANGLE),
]
TEXT_PRIMARY, TEXT_SECONDARY, GRID = "0B0B0B", "52514E", "E4E3DF"
FONT = "Arial"


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def style_text(font, size: float, color: str) -> None:
    font.name, font.size, font.color.rgb = FONT, Pt(size), rgb(color)


def main() -> None:
    df = pd.read_csv(VAL / "nhe_benchmark.csv")
    years = [str(y) for y in df["year"]]
    prelim = [bool(p) for p in df["preliminary"]]

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(7.5), Inches(4.2)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    data = CategoryChartData()
    data.categories = years
    for col, label, *_ in SERIES:
        data.add_series(label, [round(v / 1000, 2) for v in df[col]], number_format='0.00')

    gf = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS, Inches(0.1), Inches(0.1), Inches(7.3), Inches(4.0), data
    )
    chart = gf.chart
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    style_text(chart.legend.font, 9, TEXT_PRIMARY)

    va = chart.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 0, 20, 2.5
    va.has_title = True
    va.axis_title.text_frame.text = "Dental spending ($ billions)"
    style_text(va.axis_title.text_frame.paragraphs[0].runs[0].font, 9, TEXT_PRIMARY)
    va.axis_title.text_frame.paragraphs[0].runs[0].font.bold = False
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "0.0", False
    style_text(va.tick_labels.font, 9, TEXT_SECONDARY)
    va.major_gridlines.format.line.color.rgb = rgb(GRID)
    va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()

    ca = chart.category_axis
    style_text(ca.tick_labels.font, 9, TEXT_SECONDARY)
    ca.format.line.color.rgb = rgb(TEXT_SECONDARY)
    ca.has_major_gridlines = False

    for s, (col, _, color, marker) in zip(chart.plots[0].series, SERIES):
        s.smooth = False
        s.format.line.color.rgb = rgb(color)
        s.format.line.width = Pt(2)
        s.marker.style, s.marker.size = marker, 7
        s.marker.format.fill.solid()
        s.marker.format.fill.fore_color.rgb = rgb(color)
        s.marker.format.line.color.rgb = rgb("FFFFFF")

        for i, is_prelim in enumerate(prelim):
            if not is_prelim:
                continue
            # A point's format styles the segment leading into it: dash it and
            # draw the preliminary marker hollow, as in the matplotlib figure.
            pt = s.points[i]
            pt.format.line.color.rgb = rgb(color)
            pt.format.line.width = Pt(2)
            pt.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            pt.marker.style, pt.marker.size = marker, 7
            pt.marker.format.fill.solid()
            pt.marker.format.fill.fore_color.rgb = rgb("FFFFFF")
            pt.marker.format.line.color.rgb = rgb(color)
            pt.marker.format.line.width = Pt(1.5)

        if col == "chomp_50dc":
            # Mapped total as a share of NHE, under each mapped point.
            for i, pct in enumerate(df["chomp_pct_of_nhe"]):
                d = s.points[i].data_label
                d.position = XL_LABEL_POSITION.BELOW
                d.text_frame.text = f"{pct:.0f}%"
                style_text(d.text_frame.paragraphs[0].runs[0].font, 8, TEXT_SECONDARY)

    out = VAL / "fig_nhe_benchmark.pptx"
    prs.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
