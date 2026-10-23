#!/usr/bin/env python3
"""Figure 2 as an editable Excel chart: NHE vs. HHS release vs. mapped total.

Same data and styling as scripts/plot_nhe.py, written as a native Excel line chart
next to its data table. Edit the numbers in the table and the chart, the
percentage column, and the percentage labels under the mapped line all update.

openpyxl cannot link a data label to a cell, so the chart XML is patched after
saving: each mapped-series label points at its row in the Label column.

Reads docs/validation/nhe_benchmark.csv, writes docs/validation/fig_nhe_benchmark.xlsx.
"""

from __future__ import annotations

import re
import shutil
import tempfile
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.label import DataLabel, DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.marker import DataPoint, Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent
VAL = REPO_ROOT / "docs" / "validation"
SHEET = "Figure 2"

# Same colors and markers as plot_nhe.py. (csv column, header, hex, marker)
SERIES = [
    ("medicaid_chip", "NHE Medicaid + CHIP dental services ($B)", "2A78D6", "circle"),
    ("release_dental", "HHS release, all dental codes ($B)", "EB6834", "square"),
    ("chomp_50dc", "Mapped to 50 states + DC ($B)", "1BAF7A", "triangle"),
]
TEXT_PRIMARY, TEXT_SECONDARY, GRID = "0B0B0B", "52514E", "E4E3DF"
FONT = "Arial"
LINE_W = 25400  # 2 pt in EMU


def text_props(size_pt: float, color: str) -> RichText:
    cp = CharacterProperties(latin=None, sz=int(size_pt * 100), solidFill=color)
    cp.latin = None
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def marker(symbol: str, fill: str, edge: str, edge_w: int = 9525) -> Marker:
    m = Marker(symbol=symbol, size=7)
    m.graphicalProperties = GraphicalProperties(solidFill=fill)
    m.graphicalProperties.line = LineProperties(solidFill=edge, w=edge_w)
    return m


def build(out: Path) -> list[str]:
    df = pd.read_csv(VAL / "nhe_benchmark.csv")
    n = len(df)
    last = n + 1  # last data row

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    headers = ["Year"] + [h for _, h, *_ in SERIES] + ["Mapped, % of NHE", "Label"]
    ws.append(headers)
    for _, r in df.iterrows():
        row = ws.max_row + 1
        ws.append([str(int(r["year"])), *(round(r[c] / 1000, 3) for c, *_ in SERIES),
                   f"=D{row}/B{row}", f'=TEXT(E{row},"0%")'])

    thin = Side(style="thin", color="999999")
    for c in range(1, len(headers) + 1):
        h = ws.cell(row=1, column=c)
        h.font = Font(name=FONT, bold=True)
        h.alignment = Alignment(wrap_text=True, vertical="top")
        h.border = Border(bottom=thin)
        for r in range(2, last + 1):
            cell = ws.cell(row=r, column=c)
            # Blue = values typed in from the source CSV, black = formulas.
            cell.font = Font(name=FONT, color="0000FF" if 2 <= c <= 4 else "000000")
            cell.number_format = {5: "0.0%"}.get(c, "0.000" if 2 <= c <= 4 else "@")
    ws.column_dimensions["A"].width = 8
    for c in range(2, 5):
        ws.column_dimensions[get_column_letter(c)].width = 18
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 10
    ws.row_dimensions[1].height = 45

    notes = [
        "Notes",
        "Blue cells are inputs from docs/validation/nhe_benchmark.csv (scripts/validate_nhe.py), in billions of dollars. "
        "Columns E and F are formulas; the chart's percentage labels read column F.",
        "NHE: CMS National Health Expenditure Accounts, type of service and source of funds, CY 1960-2024 "
        "(2024 edition, retrieved 2026-10-12), Medicaid + CHIP rows of Total Dental Services.",
        "HHS release: all D-code payments in the HHS Medicaid Provider Spending file, by calendar year of service. "
        "Mapped: payments linked to NPPES and geocoded to the 50 states + DC.",
        "2024 is preliminary (claims for Nov-Dec 2024 incomplete at release); drawn dashed with open markers.",
    ]
    for i, t in enumerate(notes):
        r = last + 2 + i
        c = ws.cell(row=r, column=1, value=t)
        c.font = Font(name=FONT, bold=(i == 0), italic=(i > 0), color="52514E")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)  # keep notes clear of the chart
        if i > 0:
            ws.row_dimensions[r].height = 42

    ch = LineChart()
    ch.title = None
    ch.height, ch.width = 9.5, 17  # cm
    ch.legend.position = "t"
    ch.legend.txPr = text_props(9, TEXT_PRIMARY)
    data = Reference(ws, min_col=2, max_col=4, min_row=1, max_row=last)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))

    ch.y_axis.title = "Dental spending ($ billions)"
    ch.y_axis.title.tx.rich.p[0].pPr = ParagraphProperties(
        defRPr=CharacterProperties(b=False, sz=900, solidFill=TEXT_PRIMARY))
    ch.y_axis.scaling.min, ch.y_axis.scaling.max, ch.y_axis.majorUnit = 0, 20, 2.5
    ch.y_axis.number_format = "0.0"
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRID, w=9525))
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.y_axis.txPr = text_props(9, TEXT_SECONDARY)
    ch.x_axis.majorGridlines = None
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=TEXT_SECONDARY, w=9525))
    ch.x_axis.txPr = text_props(9, TEXT_SECONDARY)
    ch.x_axis.delete = False
    ch.y_axis.delete = False

    prelim = [i for i, p in enumerate(df["preliminary"]) if bool(p)]
    for s, (col, _, color, sym) in zip(ch.series, SERIES):
        s.smooth = False
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = LINE_W
        s.marker = marker(sym, color, "FFFFFF")
        for i in prelim:
            # A point's format styles the segment leading into it.
            pt = DataPoint(idx=i, marker=marker(sym, "FFFFFF", color, 19050))
            pt.graphicalProperties = GraphicalProperties(
                ln=LineProperties(solidFill=color, w=LINE_W, prstDash="dash"))
            s.dPt.append(pt)
        if col == "chomp_50dc":
            s.dLbls = DataLabelList(dLbl=[DataLabel(idx=i, dLblPos="b", showVal=True, showSerName=False,
                                                    showCatName=False, showLegendKey=False,
                                                    showPercent=False, txPr=text_props(8, TEXT_SECONDARY))
                                          for i in range(n)])
            s.dLbls.showVal = False
    ws.add_chart(ch, "H2")
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    return [f"{p:.0%}" for p in df["chomp_50dc"] / df["medicaid_chip"]]


def link_labels(path: Path, labels: list[str]) -> None:
    """Point each mapped-series data label at its cell in the Label column (F)."""
    tmp = Path(tempfile.mkdtemp())
    with zipfile.ZipFile(path) as z:
        z.extractall(tmp)
        names = z.namelist()
    chart = tmp / "xl" / "charts" / "chart1.xml"
    xml = chart.read_text(encoding="utf-8")

    def tx(m: re.Match) -> str:
        i = int(m.group(1))
        # openpyxl writes the chart namespace as the default, so no c: prefix.
        ref = (f"<tx><strRef><f>'{SHEET}'!$F${i + 2}</f><strCache><ptCount val=\"1\"/>"
               f"<pt idx=\"0\"><v>{labels[i]}</v></pt></strCache></strRef></tx>")
        return m.group(0) + ref

    xml, k = re.subn(r'<dLbl><idx val="(\d+)"\s*/>', tx, xml)
    xml = xml.replace("<a:r><a:t></a:t></a:r>", "")  # empty runs openpyxl adds to txPr
    assert k == len(labels), f"patched {k} labels, expected {len(labels)}"
    chart.write_text(xml, encoding="utf-8")
    path.unlink()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name in names:
            z.write(tmp / name, name)
    shutil.rmtree(tmp)


def main() -> None:
    out = VAL / "fig_nhe_benchmark.xlsx"
    labels = build(out)
    link_labels(out, labels)
    print("wrote", out)


if __name__ == "__main__":
    main()
